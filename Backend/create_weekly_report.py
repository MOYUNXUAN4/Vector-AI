from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\QIYUXUAN\Desktop\fyp_Backend (3)\fyp_Backend\Weekly_Report_2026-09-28_to_2026-10-02.docx"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def borders(table, color='D9D9D9'):
    tblPr = table._tbl.tblPr
    b = tblPr.first_child_found_in('w:tblBorders')
    if b is None:
        b = OxmlElement('w:tblBorders')
        tblPr.append(b)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = b.find(qn('w:' + edge))
        if e is None:
            e = OxmlElement('w:' + edge)
            b.append(e)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), color)

def set_cell(cell, text, bold=False, color='000000', size=8.5):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    for i, header in enumerate(headers):
        set_cell(table.rows[0].cells[i], header, True, 'FFFFFF', 8)
        shade(table.rows[0].cells[i], '244A64')
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value, False, '000000', 8)
            if ri % 2:
                shade(cells[i], 'F2F6F8')
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.45)
section.bottom_margin = Inches(.45)
section.left_margin = Inches(.55)
section.right_margin = Inches(.55)
styles = doc.styles
styles['Normal'].font.name = 'Arial'
styles['Normal'].font.size = Pt(9)
styles['Normal'].paragraph_format.space_after = Pt(3)
for name in ('Title', 'Heading 1', 'Heading 2'):
    styles[name].font.name = 'Arial'
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
# Word's built-in Title style can carry a blue paragraph border. Remove it so
# the report uses clean whitespace below the title.
title_ppr = styles['Title']._element.get_or_add_pPr()
title_border = title_ppr.find(qn('w:pBdr'))
if title_border is not None:
    title_ppr.remove(title_border)
styles['Title'].font.size = Pt(19)
styles['Heading 1'].font.size = Pt(12)
styles['Heading 1'].font.bold = True
styles['Heading 1'].paragraph_format.space_before = Pt(5)
styles['Heading 1'].paragraph_format.space_after = Pt(2)

p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Weekly Progress Report')
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('28 September to 2 October 2026')
r.italic = True
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(80, 80, 80)

p = doc.add_paragraph()
r = p.add_run('Main result. ')
r.bold = True
p.add_run('The project moved from a fixed geometric warning rule to a tested comparison of Time-to-Edge (TTE), LightGBM, Random Forest and GRU. LightGBM is now the main model candidate because it gives the best balance between session-level F1, precision and false alarms in development cross-validation.')

doc.add_heading('Work completed', level=1)
add_table(doc, ['Area', 'Work completed', 'Value'], [
    ['Data and labels', 'Reviewed TEST9282, TEST930 and TEST931; kept session-level splits and froze TEST929/TEST1001 for evaluation.', 'Avoids leakage and makes comparisons fair.'],
    ['Warning logic', 'Audited edge geometry, fall-event timing, post-fall warning rows and response metadata.', 'Separates real warnings from late or post-fall display states.'],
    ['Models', 'Compared the existing yellow rule, TTE, LightGBM, Random Forest and offline GRU.', 'Shows the practical accuracy/lead-time trade-off.'],
    ['Response experiment', 'Added backend analysis for attention/strong episodes, reaction time, safe return and recovery.', 'Measures whether an earlier warning helps the player.'],
], [1.05, 3.35, 2.3])

doc.add_heading('Model comparison', level=1)
doc.add_paragraph('Metrics below are session-level results on the development pool (117 sessions: 48 fall, 69 non-fall), using session-grouped cross-validation. Accuracy is not the primary metric because normal windows dominate the data.')
add_table(doc, ['Method', 'Precision', 'Recall', 'F1', 'Non-fall FAR', 'Median lead'], [
    ['TTE rule', '0.71', '1.00', '0.83', '0.29', '0.29 s'],
    ['Random Forest', '0.71', '0.71', '0.71', '0.20', '0.25 s'],
    ['GRU (0.5 s horizon)', '0.67', '0.98', '0.80', '0.33', '0.28 s'],
    ['LightGBM (selected)', '0.81', '0.90', '0.85', '0.14', '0.22 s'],
], [1.65, .75, .65, .55, 1.0, 1.0])
p = doc.add_paragraph('The historical yellow rule remains a reference only: on frozen TEST929 it had session F1 0.53 and non-fall FAR 0.20. It was not used to tune the models above.')
p.paragraph_format.space_after = Pt(2)
p = doc.add_paragraph()
r = p.add_run('Interpretation. ')
r.bold = True
p.add_run('TTE gives the highest recall and a slightly earlier warning, but raises false alarms. GRU detects more falls than the old rule but produces more non-fall alarms. LightGBM has the highest F1 and precision with the lowest false-alarm rate among the learned models, so it is the current priority.')

doc.add_heading('GRU horizon finding', level=1)
doc.add_paragraph('Longer GRU horizons increase warning time but reduce discrimination: the 0.75 s, 1.0 s and 1.5 s horizons produced median leads of about 0.47 s, 0.51 s and 0.74 s, while session F1 fell to about 0.71, 0.66 and 0.65. Earlier prediction is possible, but only with a clear precision and false-alarm cost.')

doc.add_heading('Decision and next week plan', level=1)
p = doc.add_paragraph()
r = p.add_run('Decision. ')
r.bold = True
p.add_run('Continue with LightGBM as the main offline model. Keep TTE as an interpretable high-recall reference and keep GRU as a research comparison. Do not connect a model to Unity until the response experiment confirms that players can use the earlier warning.')
add_table(doc, ['Next week', 'Planned action', 'Output'], [
    ['1', 'Collect two-level warning trials across difficulties and directions.', 'Recovery and response rates.'],
    ['2', 'Run the response analyzer on the new logs.', 'Reaction time, safe return and false alarms.'],
    ['3', 'Tune LightGBM with hard negatives and temporal smoothing.', 'Fixed threshold and warning policy.'],
    ['4', 'Validate once on frozen data using F1, FAR and lead time.', 'Final offline comparison.'],
    ['5', 'Freeze the selected LightGBM configuration and document limitations.', 'Reproducible model package.'],
], [.55, 4.25, 1.9])

doc.save(OUT)
print(OUT)
